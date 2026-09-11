from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument


@pytest.fixture
def sample_txt_path(tmp_path: Path) -> Path:
    path = tmp_path / "sample.txt"
    path.write_text(
        "Sample   text\r\n\r\n\r\nSecond paragraph with policy details.",
        encoding="utf-8",
    )
    return path


@pytest.fixture
def sample_docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.core_properties.author = "DocIntel"
    document.core_properties.title = "Employee Policy"
    document.add_heading("PTO Policy", level=1)
    document.add_paragraph("Employees receive fifteen PTO days each year.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Expense"
    table.cell(0, 1).text = "Approval"
    table.cell(1, 0).text = "Travel"
    table.cell(1, 1).text = "Manager"
    document.save(path)
    return path


@pytest.fixture
def empty_docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "empty.docx"
    DocxDocument().save(path)
    return path


@pytest.fixture
def sample_pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "sample.pdf"
    pdf = pymupdf.open()
    pdf.set_metadata(
        {
            "title": "Security Policy",
            "author": "DocIntel",
            "subject": "Portfolio fixture",
        }
    )
    page = pdf.new_page()
    page.insert_text((72, 72), "Security Policy\nEmployees must use MFA for internal systems.")
    pdf.save(path)
    pdf.close()
    return path


@pytest.fixture
def empty_pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "empty.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()
    return path
