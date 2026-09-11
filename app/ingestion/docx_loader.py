from pathlib import Path

from docx import Document as DocxDocument

from app.documents.models import Document
from app.ingestion.base import DocumentExtractionError, DocumentLoader, EmptyDocumentError
from app.ingestion.utils import build_file_document


class DocxLoader(DocumentLoader):
    def load(self, path: Path) -> Document:
        try:
            docx = DocxDocument(path)
        except Exception as exc:
            raise DocumentExtractionError(f"Could not read DOCX file '{path.name}'.") from exc

        parts: list[str] = []
        for paragraph in docx.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in docx.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    parts.append(row_text)

        text = "\n\n".join(parts)
        if not text.strip():
            raise EmptyDocumentError(f"DOCX file '{path.name}' contains no readable text.")

        core = docx.core_properties
        metadata = {
            "file_type": "docx",
            "paragraph_count": len(docx.paragraphs),
            "table_count": len(docx.tables),
            "author": core.author or None,
            "title": core.title or None,
            "subject": core.subject or None,
        }

        return build_file_document(path=path, text=text, metadata=metadata)
